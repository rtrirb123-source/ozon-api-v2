#!/usr/bin/env python3
import hashlib
import json
import os
import re
import socket
import ssl
import tempfile
from datetime import datetime, timezone
from pathlib import Path

import requests
import urllib3
from openpyxl import load_workbook

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_sku(value):
    return re.sub(r"\s+", "", clean(value)).upper()


def parse_boxes(value):
    raw = clean(value)
    if not raw:
        return [], "未分配为空"
    if any(word in raw for word in ("空盒", "盒子", "已使用")):
        return [], raw
    parts = [part.strip() for part in re.split(r"[,，、;；]", raw) if part.strip()]
    boxes = []
    for part in parts:
        if re.fullmatch(r"\d+", part):
            boxes.append(int(part))
            continue
        match = re.fullmatch(r"(\d+)\s*-\s*(\d+)", part)
        if match:
            start, end = map(int, match.groups())
            if end < start or end - start > 5000:
                return [], f"非法箱号范围：{part}"
            boxes.extend(range(start, end + 1))
            continue
        return [], f"无法解析：{part}"
    return boxes, "" if boxes else "未识别箱号"


def find_header(rows):
    for row_index, row in enumerate(rows[:30]):
        values = [clean(value) for value in row]
        if "SKU" in values and "单箱数量" in values and "未分配" in values:
            return {
                "row": row_index,
                "sku": values.index("SKU"),
                "per_box": values.index("单箱数量"),
                "unallocated": values.index("未分配"),
                "product": values.index("产品品名") if "产品品名" in values else -1,
            }
    return None


def parse_workbook(path, source_file):
    workbook = load_workbook(path, read_only=True, data_only=True)
    records, warnings = [], []
    try:
        for sheet in workbook.worksheets:
            rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
            header = find_header(rows)
            if not header:
                continue
            box_mark = ""
            for row_index, row in enumerate(rows):
                values = list(row)
                mark_index = next((i for i, value in enumerate(values) if clean(value) == "箱唛号"), -1)
                if mark_index >= 0:
                    box_mark = next((clean(value) for value in values[mark_index + 1:] if clean(value)), "")
                    continue
                if row_index <= header["row"]:
                    continue
                sku = normalize_sku(values[header["sku"]] if header["sku"] < len(values) else "")
                if not sku or sku == "SKU":
                    continue
                try:
                    per_box_qty = float(values[header["per_box"]] or 0)
                except (TypeError, ValueError):
                    per_box_qty = 0
                raw_unallocated = values[header["unallocated"]] if header["unallocated"] < len(values) else ""
                boxes, note = parse_boxes(raw_unallocated)
                if not boxes or not box_mark or per_box_qty <= 0:
                    warnings.append({
                        "source_file": source_file,
                        "source_sheet": sheet.title,
                        "source_row": row_index + 1,
                        "sku": sku,
                        "box_mark": box_mark,
                        "value": clean(raw_unallocated),
                        "note": note or "缺少箱唛或单箱数量",
                    })
                    continue
                product_name = clean(values[header["product"]]) if header["product"] >= 0 and header["product"] < len(values) else ""
                records.append({
                    "source_sheet": sheet.title,
                    "source_row": row_index + 1,
                    "box_mark": box_mark,
                    "box_numbers": clean(raw_unallocated),
                    "sku": sku,
                    "product_name": product_name,
                    "per_box_qty": per_box_qty,
                    "box_count": len(boxes),
                    "pieces": len(boxes) * per_box_qty,
                    "note": "",
                })
    finally:
        workbook.close()
    return records, warnings


def verify_certificate(base_url, expected):
    from urllib.parse import urlparse
    parsed = urlparse(base_url)
    host, port = parsed.hostname, parsed.port or 443
    context = ssl._create_unverified_context()
    with socket.create_connection((host, port), timeout=15) as raw:
        with context.wrap_socket(raw, server_hostname=host) as wrapped:
            actual = hashlib.sha256(wrapped.getpeercert(binary_form=True)).hexdigest().upper()
    expected = expected.replace(":", "").strip().upper()
    if not expected or actual != expected:
        raise RuntimeError(f"NAS证书指纹不匹配：{actual}")


class Synology:
    def __init__(self, base_url, username, password):
        self.base = base_url.rstrip("/")
        self.username = username
        self.password = password
        self.session = requests.Session()
        self.session.verify = False
        self.sid = ""

    def call(self, params, stream=False):
        params = dict(params)
        if self.sid:
            params["_sid"] = self.sid
        response = self.session.get(f"{self.base}/webapi/entry.cgi", params=params, timeout=120, stream=stream)
        response.raise_for_status()
        return response

    def login(self):
        response = self.call({
            "api": "SYNO.API.Auth", "version": "7", "method": "login",
            "account": self.username, "passwd": self.password,
            "session": "FileStation", "format": "sid",
        })
        payload = response.json()
        if not payload.get("success"):
            raise RuntimeError(f"NAS登录失败，错误码：{payload.get('error', {}).get('code')}")
        self.sid = payload["data"]["sid"]

    def list_xlsx(self, folder):
        response = self.call({
            "api": "SYNO.FileStation.List", "version": "2", "method": "list",
            "folder_path": folder, "offset": 0, "limit": 1000,
            "filetype": "file", "additional": json.dumps(["size", "time"]),
        })
        payload = response.json()
        if not payload.get("success"):
            raise RuntimeError(f"NAS目录读取失败，错误码：{payload.get('error', {}).get('code')}")
        return [item for item in payload.get("data", {}).get("files", [])
                if item.get("name", "").lower().endswith(".xlsx") and not item.get("name", "").startswith("~$")]

    def download(self, nas_path, local_path):
        response = self.call({
            "api": "SYNO.FileStation.Download", "version": "2", "method": "download",
            "path": nas_path, "mode": "download",
        }, stream=True)
        with open(local_path, "wb") as output:
            for chunk in response.iter_content(1024 * 1024):
                if chunk:
                    output.write(chunk)
        if Path(local_path).read_bytes()[:2] != b"PK":
            raise RuntimeError(f"NAS下载内容不是xlsx：{nas_path}")

    def logout(self):
        if not self.sid:
            return
        try:
            self.call({"api": "SYNO.API.Auth", "version": "7", "method": "logout", "session": "FileStation"})
        finally:
            self.sid = ""


def main():
    base_url = os.environ["NAS_BASE_URL"]
    verify_certificate(base_url, os.environ["NAS_CERT_SHA256"])
    client = Synology(base_url, os.environ["NAS_USERNAME"], os.environ["NAS_PASSWORD"])
    folder = os.environ["NAS_FOLDER"]
    all_files, warnings = [], []
    client.login()
    try:
        files = client.list_xlsx(folder)
        if not files:
            raise RuntimeError(f"NAS目录没有xlsx文件：{folder}")
        with tempfile.TemporaryDirectory(prefix="unallocated-nas-") as temp_dir:
            for item in sorted(files, key=lambda value: value["name"]):
                local_path = Path(temp_dir) / item["name"]
                client.download(item.get("path") or f"{folder}/{item['name']}", local_path)
                records, file_warnings = parse_workbook(local_path, item["name"])
                all_files.append({"source_file": item["name"], "records": records})
                warnings.extend(file_warnings)
    finally:
        client.logout()

    response = requests.post(
        os.environ["INVENTORY_IMPORT_URL"],
        headers={"content-type": "application/json", "x-inventory-sync-token": os.environ["INVENTORY_SYNC_TOKEN"]},
        json={"files": all_files}, timeout=120,
    )
    payload = response.json()
    if not response.ok or not payload.get("ok"):
        raise RuntimeError(payload.get("error") or f"库存导入失败：HTTP {response.status_code}")
    records = [row for item in all_files for row in item["records"]]
    report = {
        "at": datetime.now(timezone.utc).isoformat(),
        "file_count": len(all_files), "lines": len(records),
        "sku_count": len({row["sku"] for row in records}),
        "boxes": sum(row["box_count"] for row in records),
        "pieces": sum(row["pieces"] for row in records),
        "warning_count": len(warnings), "warnings": warnings,
        "server": payload.get("data"),
    }
    report_dir = Path("/var/lib/ozon-unallocated-sync")
    report_dir.mkdir(parents=True, exist_ok=True)
    (report_dir / "last-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({key: value for key, value in report.items() if key != "warnings"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
