#!/usr/bin/env python3
import importlib.util
import json
import os
import tempfile
from datetime import datetime, timezone
from pathlib import Path

import requests
from openpyxl import load_workbook

BASE_SCRIPT = "/opt/ozon-api-v2/scripts/nas_unallocated_cloud_sync.py"
spec = importlib.util.spec_from_file_location("nas_common", BASE_SCRIPT)
common = importlib.util.module_from_spec(spec)
spec.loader.exec_module(common)


def find_header(rows):
    for index, row in enumerate(rows[:30]):
        values = [common.clean(value) for value in row]
        if "SKU" in values and "总数" in values:
            return index, values.index("SKU"), values.index("总数")
    return None


def parse_workbook(path, source_file):
    book = load_workbook(path, read_only=True, data_only=True)
    records, warnings = [], []
    try:
        for sheet in book.worksheets:
            rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
            header = find_header(rows)
            if not header:
                if any(any(common.clean(value) for value in row) for row in rows):
                    warnings.append({"source_file": source_file, "sheet": sheet.title, "note": "未找到SKU和总数表头"})
                continue
            header_row, sku_col, total_col = header
            box_mark = ""
            for row_number, row in enumerate(rows, start=1):
                values = list(row)
                mark_index = next((i for i, value in enumerate(values) if common.clean(value) == "箱唛号"), -1)
                if mark_index >= 0:
                    box_mark = next((common.clean(value) for value in values[mark_index + 1:] if common.clean(value)), "")
                    continue
                if row_number <= header_row + 1:
                    continue
                sku = common.normalize_sku(row[sku_col] if sku_col < len(row) else "")
                raw_total = row[total_col] if total_col < len(row) else None
                if not sku:
                    continue
                try:
                    quantity = float(raw_total or 0)
                except (TypeError, ValueError):
                    warnings.append({"source_file": source_file, "sheet": sheet.title, "row": row_number, "sku": sku, "value": common.clean(raw_total), "note": "总数不是数字"})
                    continue
                if quantity <= 0:
                    continue
                records.append({
                    "source_sheet": sheet.title,
                    "source_row": row_number,
                    "box_mark": box_mark,
                    "sku": sku,
                    "quantity": quantity,
                })
    finally:
        book.close()
    return records, warnings


def main():
    common.verify_certificate(os.environ["NAS_BASE_URL"], os.environ["NAS_CERT_SHA256"])
    client = common.Synology(os.environ["NAS_BASE_URL"], os.environ["NAS_USERNAME"], os.environ["NAS_PASSWORD"])
    folder = os.environ.get("NAS_FIRST_LEG_FOLDER", "/公司资料/俄罗斯/俄罗斯在途库存")
    files_payload, warnings = [], []
    client.login()
    try:
        files = client.list_xlsx(folder)
        if not files:
            raise RuntimeError(f"NAS目录没有xlsx文件：{folder}")
        with tempfile.TemporaryDirectory(prefix="first-leg-nas-") as temp_dir:
            for item in sorted(files, key=lambda value: value["name"]):
                target = Path(temp_dir) / item["name"]
                client.download(item.get("path") or f"{folder}/{item['name']}", target)
                records, file_warnings = parse_workbook(target, item["name"])
                files_payload.append({"source_file": item["name"], "records": records})
                warnings.extend(file_warnings)
    finally:
        client.logout()

    response = requests.post(
        os.environ.get("FIRST_LEG_IMPORT_URL", "http://127.0.0.1:3000/api/inventory/first-leg-transit/import"),
        headers={"content-type": "application/json", "x-inventory-sync-token": os.environ["INVENTORY_SYNC_TOKEN"]},
        json={"files": files_payload}, timeout=120,
    )
    payload = response.json()
    if not response.ok or not payload.get("ok"):
        raise RuntimeError(payload.get("error") or f"头程在途导入失败：HTTP {response.status_code}")
    server = payload.get("data") or {}
    report = {"at": datetime.now(timezone.utc).isoformat(), "warning_count": len(warnings), "warnings": warnings, "server": server}
    report_dir = Path("/var/lib/ozon-first-leg-sync")
    report_dir.mkdir(parents=True, exist_ok=True)
    (report_dir / "last-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"at": report["at"], "warning_count": len(warnings), "server": server}, ensure_ascii=False))


if __name__ == "__main__":
    main()
